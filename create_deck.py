"""Generate the AdPulse presentation deck (.pptx)."""

import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Paths ---
BASE = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE, "deck_assets")
os.makedirs(IMG_DIR, exist_ok=True)
OUTPUT = os.path.join(BASE, "AdPulse_Presentation.pptx")

# --- Colors ---
NAVY = RGBColor(0x1A, 0x1F, 0x36)
TEAL = RGBColor(0x08, 0x91, 0xB2)
CORAL = RGBColor(0xFF, 0x6B, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF8, 0xFA, 0xFC)
DARK_TEXT = RGBColor(0x1E, 0x29, 0x3B)
MID_GRAY = RGBColor(0x64, 0x74, 0x8B)
ICE_BLUE = RGBColor(0xCA, 0xDC, 0xFC)

HEX_NAVY = "#1A1F36"
HEX_TEAL = "#0891B2"
HEX_CORAL = "#FF6B6B"
HEX_WHITE = "#FFFFFF"
HEX_DARK = "#1E293B"
HEX_MID = "#64748B"

# --- Chart generation ---

def load_data():
    sys.path.insert(0, BASE)
    from data_utils import load_superbowl_data
    from analysis import SentimentAnalyzer, EngagementAnalyzer
    df = load_superbowl_data()
    return df

def make_feature_chart(df):
    from analysis import EngagementAnalyzer
    eng = EngagementAnalyzer(df)
    impact = eng.feature_impact()
    impact = impact.sort_values("views_lift")

    fig, ax = plt.subplots(figsize=(8, 4))
    colors = [HEX_TEAL if v >= 0 else HEX_CORAL for v in impact["views_lift"]]
    bars = ax.barh(impact["feature"], impact["views_lift"], color=colors, height=0.6)
    ax.set_xlabel("Views Lift (%)", fontsize=11, color=HEX_MID)
    ax.set_title("Impact of Ad Features on Views", fontsize=14, fontweight="bold", color=HEX_DARK, pad=15)
    ax.axvline(0, color="#CBD5E1", linewidth=0.8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["bottom"].set_color("#E2E8F0")
    ax.spines["left"].set_color("#E2E8F0")
    ax.tick_params(colors=HEX_MID)
    for bar, val in zip(bars, impact["views_lift"]):
        ax.text(bar.get_width() + (2 if val >= 0 else -2), bar.get_y() + bar.get_height()/2,
                f"{val:+.0f}%", va="center", ha="left" if val >= 0 else "right",
                fontsize=10, color=HEX_DARK)
    plt.tight_layout()
    path = os.path.join(IMG_DIR, "feature_impact.png")
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close()
    return path

def make_sentiment_chart(df):
    from analysis import SentimentAnalyzer
    sa = SentimentAnalyzer("vader")
    descs = df[df["description"].notna()]["description"].tolist()
    results = sa.analyze(descs)
    counts = results["sentiment"].value_counts()

    fig, ax = plt.subplots(figsize=(5, 5))
    colors_map = {"positive": "#2ca02c", "neutral": "#94a3b8", "negative": "#ef4444"}
    labels = counts.index.tolist()
    vals = counts.values.tolist()
    cols = [colors_map.get(l, "#999") for l in labels]
    wedges, texts, autotexts = ax.pie(
        vals, labels=[l.title() for l in labels], autopct="%1.0f%%",
        colors=cols, startangle=90, textprops={"fontsize": 12}
    )
    for t in autotexts:
        t.set_fontweight("bold")
        t.set_color("white")
    ax.set_title("Sentiment of Ad Descriptions", fontsize=14, fontweight="bold", color=HEX_DARK, pad=15)
    plt.tight_layout()
    path = os.path.join(IMG_DIR, "sentiment_pie.png")
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close()
    return path

def make_brand_chart(df):
    from analysis import EngagementAnalyzer
    brand_sum = EngagementAnalyzer(df).brand_summary().head(6)
    brand_sum = brand_sum.sort_values("avg_views")

    fig, ax = plt.subplots(figsize=(8, 3.5))
    ax.barh(brand_sum["brand"], brand_sum["avg_views"] / 1e6, color=HEX_TEAL, height=0.55)
    ax.set_xlabel("Average Views (Millions)", fontsize=11, color=HEX_MID)
    ax.set_title("Top Brands by Average Views", fontsize=14, fontweight="bold", color=HEX_DARK, pad=15)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["bottom"].set_color("#E2E8F0")
    ax.spines["left"].set_color("#E2E8F0")
    ax.tick_params(colors=HEX_MID)
    for i, (_, row) in enumerate(brand_sum.iterrows()):
        ax.text(row["avg_views"]/1e6 + 0.1, i, f"{row['avg_views']/1e6:.1f}M",
                va="center", fontsize=10, color=HEX_DARK)
    plt.tight_layout()
    path = os.path.join(IMG_DIR, "brand_views.png")
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close()
    return path

def make_pipeline_diagram():
    """Create a simple technical pipeline diagram."""
    fig, ax = plt.subplots(figsize=(9, 2.8))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 3)
    ax.axis("off")

    boxes = [
        (0.3, 1, "Data\nCollection", HEX_TEAL),
        (2.3, 1, "Preprocessing\n& Cleaning", "#0E7490"),
        (4.3, 1, "AI Analysis\n(NLP + Stats)", HEX_CORAL),
        (6.3, 1, "Visualization\n& Scoring", "#8B5CF6"),
        (8.3, 1, "Interactive\nDashboard", "#059669"),
    ]
    for bx, by, label, color in boxes:
        rect = plt.Rectangle((bx, by), 1.6, 1.2, facecolor=color, edgecolor="none", alpha=0.9, zorder=2)
        ax.add_patch(rect)
        ax.text(bx + 0.8, by + 0.6, label, ha="center", va="center",
                fontsize=9, fontweight="bold", color="white", zorder=3)

    for i in range(len(boxes) - 1):
        x1 = boxes[i][0] + 1.6
        x2 = boxes[i+1][0]
        y = 1.6
        ax.annotate("", xy=(x2, y), xytext=(x1, y),
                    arrowprops=dict(arrowstyle="->", color="#94A3B8", lw=2))

    sub_items = [
        (0.3, 0.3, "FiveThirtyEight\n247 ads, 10 brands"),
        (2.3, 0.3, "NLTK tokenization\nLemmatization"),
        (4.3, 0.3, "VADER / BERT\nLDA Topics"),
        (6.3, 0.3, "Plotly charts\nComposite scoring"),
        (8.3, 0.3, "Streamlit\nReal-time filters"),
    ]
    for sx, sy, label in sub_items:
        ax.text(sx + 0.8, sy + 0.2, label, ha="center", va="center",
                fontsize=7, color=HEX_MID, style="italic")

    plt.tight_layout()
    path = os.path.join(IMG_DIR, "pipeline.png")
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


# --- Slide helpers ---

def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, prs, x, y, w, h, fill_color, transparency=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape

def add_text_box(slide, text, x, y, w, h, font_size=14, color=DARK_TEXT,
                 bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri",
                 anchor=MSO_ANCHOR.TOP):
    txBox = slide.shapes.add_textbox(x, y, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    tf.auto_size = None
    return txBox

def add_bullet_list(slide, items, x, y, w, h, font_size=14, color=DARK_TEXT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(x, y, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.space_after = Pt(6)
        p.level = 0
        pPr = p._pPr
        if pPr is None:
            from pptx.oxml.ns import qn
            pPr = p._p.get_or_add_pPr()
        from pptx.oxml.ns import qn
        buNone = pPr.findall(qn("a:buNone"))
        for bn in buNone:
            pPr.remove(bn)
        from lxml import etree
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", "\u2022")
    return txBox

def add_stat_card(slide, number, label, x, y, w=Inches(2), h=Inches(1.2),
                  num_color=TEAL, label_color=MID_GRAY):
    add_shape(slide, None, x, y, w, h, LIGHT_GRAY)
    add_text_box(slide, number, x + Inches(0.15), y + Inches(0.1),
                 w - Inches(0.3), Inches(0.65),
                 font_size=28, color=num_color, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, label, x + Inches(0.15), y + Inches(0.7),
                 w - Inches(0.3), Inches(0.4),
                 font_size=11, color=label_color, alignment=PP_ALIGN.CENTER)


# --- Build presentation ---

def build_deck():
    df = load_data()

    # Generate chart images
    feat_chart = make_feature_chart(df)
    sent_chart = make_sentiment_chart(df)
    brand_chart = make_brand_chart(df)
    pipeline_img = make_pipeline_diagram()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ═══════════════════════════════════════
    # SLIDE 1: Title
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_bg(slide, NAVY)
    # Accent bar at top
    add_shape(slide, prs, Inches(0), Inches(0), prs.slide_width, Inches(0.06), TEAL)

    add_text_box(slide, "AdPulse", Inches(1), Inches(1.8), Inches(11), Inches(1.2),
                 font_size=54, color=WHITE, bold=True, font_name="Calibri",
                 alignment=PP_ALIGN.CENTER)
    add_text_box(slide, "AI-Powered Ad Campaign Performance Analyzer",
                 Inches(1), Inches(3.0), Inches(11), Inches(0.8),
                 font_size=22, color=ICE_BLUE, alignment=PP_ALIGN.CENTER)

    # Thin line separator
    add_shape(slide, prs, Inches(4.5), Inches(4.0), Inches(4.3), Inches(0.02), TEAL)

    add_text_box(slide, "MSIS 521 \u2014 IT and Marketing in the New Economy",
                 Inches(1), Inches(4.4), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, "Winter Quarter 2025  |  University of Washington",
                 Inches(1), Inches(4.9), Inches(11), Inches(0.5),
                 font_size=13, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

    # Team placeholder
    add_text_box(slide, "Team: [Your Names Here]",
                 Inches(1), Inches(5.8), Inches(11), Inches(0.5),
                 font_size=13, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

    # ═══════════════════════════════════════
    # SLIDE 2: Context / The Problem
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "The Problem", Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "Why brands need better ad performance analytics",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Left column - stats
    add_stat_card(slide, "$7M+", "Cost per 30-sec\nSuper Bowl ad",
                  Inches(0.8), Inches(2.2), Inches(3.5), Inches(1.4), CORAL)
    add_stat_card(slide, "113M+", "Average Super Bowl\nviewers",
                  Inches(0.8), Inches(3.9), Inches(3.5), Inches(1.4), TEAL)
    add_stat_card(slide, "50+", "Ads per game,\nlimited attention",
                  Inches(0.8), Inches(5.6), Inches(3.5), Inches(1.4), RGBColor(0x8B, 0x5C, 0xF6))

    # Right column - problem description
    add_bullet_list(slide, [
        "Brands invest millions but rely on subjective measures of ad success",
        "Traditional metrics (TV ratings, reach) miss the digital conversation",
        "No unified tool to analyze engagement, sentiment, and themes at scale",
        "Marketers need data-driven insights to optimize future campaigns",
    ], Inches(5.0), Inches(2.3), Inches(7.5), Inches(4.5), font_size=16, color=DARK_TEXT)

    # ═══════════════════════════════════════
    # SLIDE 3: Our Solution
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "Our Solution: AdPulse", Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "An AI-augmented tool for ad campaign performance analysis",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Feature cards in 2x3 grid
    features = [
        ("Campaign Dashboard", "KPI metrics, views over time,\nbrand comparison charts"),
        ("Ad Performance Ranking", "Composite scoring across views,\nlikes, and comments"),
        ("Feature Impact Analysis", "Which ad characteristics\ndrive engagement"),
        ("Sentiment Analysis", "VADER & BERT NLP on\nad descriptions and comments"),
        ("Topic Discovery", "LDA topic modeling to find\naudience discussion themes"),
        ("Search Trends", "Google Trends integration\nfor brand search interest"),
    ]
    for i, (title, desc) in enumerate(features):
        col = i % 3
        row = i // 3
        fx = Inches(0.8) + col * Inches(4.1)
        fy = Inches(2.2) + row * Inches(2.3)
        # Card background
        add_shape(slide, prs, fx, fy, Inches(3.8), Inches(2.0), LIGHT_GRAY)
        # Left accent
        add_shape(slide, prs, fx, fy, Inches(0.06), Inches(2.0), TEAL if col != 1 else CORAL)
        # Card text
        add_text_box(slide, title, fx + Inches(0.25), fy + Inches(0.2),
                     Inches(3.3), Inches(0.45), font_size=15, color=DARK_TEXT, bold=True)
        add_text_box(slide, desc, fx + Inches(0.25), fy + Inches(0.7),
                     Inches(3.3), Inches(1.1), font_size=12, color=MID_GRAY)

    # ═══════════════════════════════════════
    # SLIDE 4: Data
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "The Data", Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "Real-world dataset from FiveThirtyEight / TidyTuesday",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Stats row
    data_stats = [
        ("247", "Super Bowl Ads"),
        ("10", "Major Brands"),
        ("21", "Years (2000-2020)"),
        ("207", "Video Descriptions"),
        ("25", "Data Columns"),
    ]
    for i, (num, label) in enumerate(data_stats):
        sx = Inches(0.6) + i * Inches(2.5)
        add_stat_card(slide, num, label, sx, Inches(2.0), Inches(2.2), Inches(1.3))

    # Data details
    add_text_box(slide, "Dataset Features", Inches(0.8), Inches(3.8), Inches(5), Inches(0.5),
                 font_size=18, color=DARK_TEXT, bold=True)
    add_bullet_list(slide, [
        "Ad characteristics: funny, celebrity, patriotic, danger, animals, use_sex, show_product_quickly",
        "YouTube engagement: views, likes, dislikes, comments, favorites",
        "Video metadata: title, description, publish date, channel, thumbnail",
        "Source: FiveThirtyEight staffers hand-coded every ad from video",
    ], Inches(0.8), Inches(4.3), Inches(11.5), Inches(2.8), font_size=14, color=DARK_TEXT)

    # Brand chart
    slide.shapes.add_picture(brand_chart, Inches(8.2), Inches(3.6), Inches(4.8), Inches(2.8))

    # ═══════════════════════════════════════
    # SLIDE 5: Technical Architecture
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "Technical Architecture", Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "End-to-end AI pipeline built in Python",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Pipeline diagram
    slide.shapes.add_picture(pipeline_img, Inches(0.5), Inches(2.0), Inches(12.3), Inches(3.2))

    # Tech stack
    add_text_box(slide, "Technology Stack", Inches(0.8), Inches(5.4), Inches(3), Inches(0.4),
                 font_size=16, color=DARK_TEXT, bold=True)
    stack_items = [
        ("AI/NLP", "VADER sentiment, BERT (optional), LDA topic modeling, NLTK preprocessing"),
        ("Data", "pandas, scikit-learn, NumPy, pytrends (Google Trends API)"),
        ("Visualization", "Plotly interactive charts, WordCloud, Matplotlib"),
        ("Interface", "Streamlit web app with real-time filtering and CSV upload"),
    ]
    for i, (cat, desc) in enumerate(stack_items):
        ty = Inches(5.85) + i * Inches(0.35)
        add_text_box(slide, f"{cat}:  {desc}", Inches(0.8), ty, Inches(12), Inches(0.35),
                     font_size=12, color=MID_GRAY)

    # ═══════════════════════════════════════
    # SLIDE 6: Demo - Feature Impact
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, CORAL)

    add_text_box(slide, "What Makes an Ad Successful?", Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "Feature impact analysis from 247 Super Bowl ads",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    slide.shapes.add_picture(feat_chart, Inches(0.5), Inches(2.0), Inches(8.5), Inches(4.5))

    # Key insight callout
    add_shape(slide, prs, Inches(9.3), Inches(2.5), Inches(3.5), Inches(3.5), LIGHT_GRAY)
    add_shape(slide, prs, Inches(9.3), Inches(2.5), Inches(0.06), Inches(3.5), TEAL)
    add_text_box(slide, "Key Insights", Inches(9.55), Inches(2.7), Inches(3.1), Inches(0.4),
                 font_size=16, color=DARK_TEXT, bold=True)
    add_bullet_list(slide, [
        "Show Product Quickly: +70% views lift",
        "Patriotic themes: +33% lift",
        "Funny ads: +12% lift",
        "Celebrity ads actually get fewer views (-54%)",
    ], Inches(9.55), Inches(3.2), Inches(3.1), Inches(2.5), font_size=12, color=DARK_TEXT)

    # ═══════════════════════════════════════
    # SLIDE 7: Demo - Sentiment & Topics
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, CORAL)

    add_text_box(slide, "NLP Analysis: Sentiment & Topics", Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "AI-powered text analysis of ad descriptions",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Sentiment chart
    slide.shapes.add_picture(sent_chart, Inches(0.5), Inches(2.0), Inches(5.0), Inches(4.5))

    # Topics section
    add_shape(slide, prs, Inches(6.0), Inches(2.0), Inches(6.8), Inches(4.5), LIGHT_GRAY)
    add_shape(slide, prs, Inches(6.0), Inches(2.0), Inches(0.06), Inches(4.5), RGBColor(0x8B, 0x5C, 0xF6))

    add_text_box(slide, "LDA Topic Discovery", Inches(6.3), Inches(2.2), Inches(6.2), Inches(0.4),
                 font_size=18, color=DARK_TEXT, bold=True)
    add_text_box(slide, "5 topics extracted from 207 ad descriptions:",
                 Inches(6.3), Inches(2.7), Inches(6.2), Inches(0.4),
                 font_size=13, color=MID_GRAY)

    topics = [
        "Topic 1: Hyundai campaigns (hyundai, bowl, super, world)",
        "Topic 2: Bud Light / Budweiser (commercial, light, bud, budweiser)",
        "Topic 3: Recent Super Bowl (bowl, super, commercial, 2020)",
        "Topic 4: Doritos / snack brands (doritos, chip, crash, super)",
        "Topic 5: NFL / game themes (nfl, game, football, team)",
    ]
    add_bullet_list(slide, topics, Inches(6.3), Inches(3.2), Inches(6.2), Inches(2.8),
                    font_size=13, color=DARK_TEXT)

    add_text_box(slide, "Methods: VADER (dictionary-based) + optional BERT transformer",
                 Inches(6.3), Inches(6.0), Inches(6.2), Inches(0.4),
                 font_size=11, color=MID_GRAY)

    # ═══════════════════════════════════════
    # SLIDE 8: Output Validation
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "Output Validation & Sanity Checking",
                 Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "Ensuring our AI outputs are reliable",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Validation cards
    val_items = [
        ("Sentiment Distribution Check",
         "Verifies that sentiment results are not\n"
         "dominated by a single class (>95% threshold).\n"
         "Flags unreliable outputs automatically.",
         "Result: PASS \u2014 75% positive, 15% neutral, 10% negative",
         TEAL),
        ("Known-Sample Accuracy",
         "Runs the model against 5 known-sentiment\n"
         "examples (clear positive, negative, neutral).\n"
         "Reports accuracy as a sanity check.",
         "Result: 4/5 correct (80% accuracy)",
         RGBColor(0x05, 0x96, 0x69)),
        ("Topic Distinctness Score",
         "Measures cosine distance between topic-word\n"
         "distributions. Higher = more distinct topics.\n"
         "Alerts when topics overlap excessively.",
         "Result: 0.312 \u2014 topics are well-separated",
         RGBColor(0x8B, 0x5C, 0xF6)),
    ]
    for i, (title, desc, result, accent) in enumerate(val_items):
        vx = Inches(0.6) + i * Inches(4.2)
        vy = Inches(2.2)
        add_shape(slide, prs, vx, vy, Inches(3.9), Inches(4.5), LIGHT_GRAY)
        add_shape(slide, prs, vx, vy, Inches(0.06), Inches(4.5), accent)
        add_text_box(slide, title, vx + Inches(0.25), vy + Inches(0.2),
                     Inches(3.4), Inches(0.5), font_size=16, color=DARK_TEXT, bold=True)
        add_text_box(slide, desc, vx + Inches(0.25), vy + Inches(0.8),
                     Inches(3.4), Inches(2.0), font_size=13, color=MID_GRAY)
        # Result box
        add_shape(slide, prs, vx + Inches(0.15), vy + Inches(3.2), Inches(3.6), Inches(1.0), WHITE)
        add_text_box(slide, result, vx + Inches(0.3), vy + Inches(3.3),
                     Inches(3.3), Inches(0.8), font_size=12, color=accent, bold=True)

    # ═══════════════════════════════════════
    # SLIDE 9: Key Findings
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "Key Findings", Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
                 font_size=36, color=DARK_TEXT, bold=True)

    findings = [
        ("Show Product Quickly wins",
         "Ads that show the product early get 70% more views on average. "
         "Clarity beats cleverness for engagement."),
        ("Celebrity doesn't guarantee success",
         "Contrary to conventional wisdom, celebrity ads averaged 54% fewer views. "
         "The product and story matter more."),
        ("Sentiment is overwhelmingly positive",
         "75% of ad descriptions carry positive sentiment. "
         "Brands craft optimistic narratives, but negative ads are rare and underexplored."),
        ("Doritos dominates engagement",
         "Doritos leads all brands with 7.9M average views per ad, "
         "largely driven by their crowd-sourced 'Crash the Super Bowl' contests."),
        ("Ad strategies shift over time",
         "Humor and sex appeal in ads have declined since 2010, "
         "while product-focused and patriotic ads have increased."),
    ]
    for i, (title, desc) in enumerate(findings):
        fy = Inches(1.5) + i * Inches(1.15)
        add_shape(slide, prs, Inches(0.8), fy, Inches(0.06), Inches(0.95), TEAL if i % 2 == 0 else CORAL)
        add_text_box(slide, title, Inches(1.1), fy + Inches(0.05), Inches(3.5), Inches(0.35),
                     font_size=15, color=DARK_TEXT, bold=True)
        add_text_box(slide, desc, Inches(4.8), fy + Inches(0.05), Inches(8), Inches(0.85),
                     font_size=13, color=MID_GRAY)

    # ═══════════════════════════════════════
    # SLIDE 10: Extensibility & Impact
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, WHITE)
    add_shape(slide, prs, Inches(0), Inches(0), Inches(0.06), prs.slide_height, TEAL)

    add_text_box(slide, "Beyond Super Bowl: Extensibility & Impact",
                 Inches(0.8), Inches(0.5), Inches(11), Inches(0.6),
                 font_size=36, color=DARK_TEXT, bold=True)
    add_text_box(slide, "AdPulse works with any ad campaign, not just Super Bowl",
                 Inches(0.8), Inches(1.15), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY)

    # Left: Extensibility
    add_text_box(slide, "Extensible by Design", Inches(0.8), Inches(2.2), Inches(5.5), Inches(0.5),
                 font_size=20, color=DARK_TEXT, bold=True)
    add_bullet_list(slide, [
        "Upload any campaign CSV via the built-in Upload page",
        "Automatic column detection for engagement metrics and text fields",
        "Sentiment and topic analysis on any text data",
        "Google Trends integration works for any brand name",
        "YouTube comment fetcher works for any video URL",
    ], Inches(0.8), Inches(2.8), Inches(5.5), Inches(3.5), font_size=14, color=DARK_TEXT)

    # Right: Impact
    add_text_box(slide, "Who Would Use This?", Inches(7.0), Inches(2.2), Inches(5.5), Inches(0.5),
                 font_size=20, color=DARK_TEXT, bold=True)

    users = [
        ("Brand Managers", "Optimize ad spend by understanding what features drive engagement"),
        ("Marketing Agencies", "Benchmark client ads against competitors in the same category"),
        ("Media Buyers", "Data-driven decisions on ad placement and creative direction"),
        ("Marketing Researchers", "Academic analysis of advertising trends and audience reception"),
    ]
    for i, (role, desc) in enumerate(users):
        uy = Inches(2.9) + i * Inches(1.05)
        add_shape(slide, prs, Inches(7.0), uy, Inches(5.8), Inches(0.85), LIGHT_GRAY)
        add_shape(slide, prs, Inches(7.0), uy, Inches(0.06), Inches(0.85), TEAL)
        add_text_box(slide, role, Inches(7.25), uy + Inches(0.05), Inches(5.3), Inches(0.35),
                     font_size=14, color=DARK_TEXT, bold=True)
        add_text_box(slide, desc, Inches(7.25), uy + Inches(0.4), Inches(5.3), Inches(0.4),
                     font_size=12, color=MID_GRAY)

    # ═══════════════════════════════════════
    # SLIDE 11: Live Demo
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    add_shape(slide, prs, Inches(0), Inches(0), prs.slide_width, Inches(0.06), TEAL)

    add_text_box(slide, "Live Demo", Inches(1), Inches(2.2), Inches(11), Inches(1.0),
                 font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, "AdPulse  \u2014  Interactive Streamlit Dashboard",
                 Inches(1), Inches(3.4), Inches(11), Inches(0.6),
                 font_size=20, color=ICE_BLUE, alignment=PP_ALIGN.CENTER)

    add_shape(slide, prs, Inches(4.5), Inches(4.3), Inches(4.3), Inches(0.02), TEAL)

    add_text_box(slide, "streamlit run app.py",
                 Inches(1), Inches(4.8), Inches(11), Inches(0.5),
                 font_size=16, color=MID_GRAY, alignment=PP_ALIGN.CENTER, font_name="Consolas")

    # ═══════════════════════════════════════
    # SLIDE 12: Thank You / Q&A
    # ═══════════════════════════════════════
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    add_shape(slide, prs, Inches(0), Inches(0), prs.slide_width, Inches(0.06), TEAL)

    add_text_box(slide, "Thank You", Inches(1), Inches(2.0), Inches(11), Inches(1.2),
                 font_size=48, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, "Questions?",
                 Inches(1), Inches(3.3), Inches(11), Inches(0.6),
                 font_size=22, color=ICE_BLUE, alignment=PP_ALIGN.CENTER)

    add_shape(slide, prs, Inches(4.5), Inches(4.2), Inches(4.3), Inches(0.02), TEAL)

    add_text_box(slide, "AdPulse  |  AI-Powered Ad Campaign Performance Analyzer",
                 Inches(1), Inches(4.7), Inches(11), Inches(0.5),
                 font_size=14, color=MID_GRAY, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, "MSIS 521  \u2014  Winter 2025  \u2014  University of Washington",
                 Inches(1), Inches(5.2), Inches(11), Inches(0.5),
                 font_size=13, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

    # Save
    prs.save(OUTPUT)
    print(f"Presentation saved to: {OUTPUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
